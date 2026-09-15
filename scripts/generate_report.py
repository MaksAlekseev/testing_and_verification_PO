from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "doc" / "Практическая_работа_1_Office_Booking.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(12)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_sep, text, fld_char_end])


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)
    section.different_first_page_header_footer = True
    add_page_number(section.footer.paragraphs[0])

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_after = Pt(0)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(14)
        style.font.bold = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def add_paragraph(document, text, align=None, first_line=True):
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(1.25 if first_line else 0)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style=None)
        paragraph.paragraph_format.left_indent = Cm(1.25)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.add_run("- " + item)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value))
    document.add_paragraph()
    return table


def add_title_page(document):
    for line in [
        "МИНОБРНАУКИ РОССИИ",
        "Федеральное государственное бюджетное образовательное учреждение высшего образования",
        "«МИРЭА - Российский технологический университет»",
        "",
        "Институт информационных технологий",
        "",
        "ОТЧЕТ",
        "по практической работе N 1",
        "по дисциплине «Тестирование и верификация ПО»",
        "",
        "Тема: «Тестирование программного продукта методом черного ящика»",
        "",
        "Проект: «Система бронирования переговорных комнат и рабочих мест в офисе»",
    ]:
        paragraph = add_paragraph(document, line, WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
        if line in ["ОТЧЕТ", "Тема: «Тестирование программного продукта методом черного ящика»"]:
            for run in paragraph.runs:
                run.bold = True

    document.add_paragraph()
    document.add_paragraph()
    add_paragraph(document, "Команда: ________________________________", WD_ALIGN_PARAGRAPH.RIGHT, False)
    add_paragraph(document, "Состав команды:", WD_ALIGN_PARAGRAPH.RIGHT, False)
    for index in range(1, 5):
        add_paragraph(document, f"{index}. ________________________________", WD_ALIGN_PARAGRAPH.RIGHT, False)
    add_paragraph(document, "Дата выполнения: 16.09.2026", WD_ALIGN_PARAGRAPH.RIGHT, False)
    document.add_paragraph()
    document.add_paragraph()
    add_paragraph(document, "Москва, 2026", WD_ALIGN_PARAGRAPH.CENTER, False)
    document.add_page_break()


def main():
    document = Document()
    configure_document(document)
    add_title_page(document)

    document.add_heading("1. Цель работы", level=1)
    add_paragraph(document, "Цель практической работы - разработать простой программный продукт, подготовить техническое задание и документацию, а также создать основу для последующего тестирования методом «черного ящика».")

    document.add_heading("2. Описание программного продукта", level=1)
    add_paragraph(document, "В рамках работы разработано настольное приложение Office Booking для бронирования переговорных комнат и рабочих мест в офисе. Пользователь выбирает ресурс, вводит сотрудника, дату, время, количество людей и цель бронирования, после чего запись отображается в таблице и может быть сохранена в CSV-файл.")
    add_paragraph(document, "Для реализации выбран стек Java 8 и Swing. Такой вариант не требует сложной инфраструктуры, подходит для учебного проекта и позволяет собрать приложение в запускаемый JAR-файл, а при наличии JDK 14+ - в EXE-образ.")

    document.add_heading("3. Техническое задание", level=1)
    document.add_heading("3.1. Введение", level=2)
    add_paragraph(document, "Программа предназначена для учета бронирования офисных ресурсов: переговорных комнат и рабочих мест. Область применения - внутренние офисные процессы малой команды или учебная демонстрация процесса тестирования.")
    document.add_heading("3.2. Основания для разработки", level=2)
    add_paragraph(document, "Основанием является практическая работа N 1 по дисциплине «Тестирование и верификация ПО». Продукт разрабатывается для последующей передачи другой команде, которая будет тестировать его методом «черного ящика».")
    document.add_heading("3.3. Назначение разработки", level=2)
    add_paragraph(document, "Назначение разработки - автоматизировать простую регистрацию броней и предоставить проверяемый программный продукт с заранее внесенными дефектами.")
    document.add_heading("3.4. Функциональные требования", level=2)
    add_bullets(document, [
        "просмотр списка доступных переговорных комнат и рабочих мест;",
        "создание брони с указанием ресурса, сотрудника, даты, времени, количества людей и цели;",
        "отображение созданных броней в таблице;",
        "отмена выбранной брони;",
        "сохранение броней в локальный CSV-файл;",
        "загрузка сохраненных броней при повторном запуске."
    ])
    document.add_heading("3.5. Требования к интерфейсу", level=2)
    add_paragraph(document, "Главное окно содержит выпадающий список ресурсов, поля ввода параметров бронирования, кнопки создания, отмены и сохранения, а также таблицу текущих броней.")
    document.add_heading("3.6. Условия эксплуатации и совместимость", level=2)
    add_paragraph(document, "Приложение рассчитано на Windows 10/11. Для запуска JAR-файла требуется Java 8 или новее. Для создания EXE-образа требуется JDK 14 или новее с утилитой jpackage.")
    document.add_heading("3.7. Критерии приемки", level=2)
    add_bullets(document, [
        "приложение запускается без аварийного завершения;",
        "пользователь может создать бронь и увидеть ее в таблице;",
        "данные можно сохранить в файл;",
        "документация содержит ТЗ, руководство пользователя и описание внесенных дефектов;",
        "в продукт внесено от 5 до 8 дефектов различной природы."
    ])

    document.add_heading("4. Руководство пользователя", level=1)
    add_paragraph(document, "Для запуска из исходников необходимо выполнить скрипт scripts/build.ps1, затем запустить команду java -jar dist/office-booking.jar. Для создания брони пользователь заполняет поля формы и нажимает кнопку «Забронировать». Для отмены записи нужно выбрать строку таблицы и нажать кнопку «Отменить выбранную бронь». Для сохранения данных используется кнопка «Сохранить».")

    document.add_heading("5. Описание внесенных дефектов", level=1)
    add_table(document,
              ["ID", "Тип", "Описание", "Условие обнаружения"],
              [
                  ["BUG-001", "Логический", "Разрешены пересекающиеся брони одного ресурса.", "Создать две брони одной комнаты на пересекающееся время."],
                  ["BUG-002", "Валидация", "Принимается время окончания меньше или равное времени начала.", "Ввести начало 14 и окончание 12."],
                  ["BUG-003", "Граничное значение", "Вместимость проверяется с ошибкой на 1 человека.", "Для рабочего места указать 2 человека."],
                  ["BUG-004", "Интерфейс / данные", "Рабочее место отображается как переговорная.", "Создать бронь рабочего места."],
                  ["BUG-005", "Логический", "Отмена удаляет первую бронь сотрудника, а не выбранную строку.", "Создать две брони на одного сотрудника и отменить вторую."],
                  ["BUG-006", "Сохранение", "Поле цели не сохраняется в CSV.", "Сохранить бронь с целью и перезапустить приложение."],
              ])

    document.add_heading("6. Примеры тест-кейсов для черного ящика", level=1)
    add_table(document,
              ["ID", "Название", "Шаги", "Ожидаемый результат", "Фактический результат", "Статус"],
              [
                  ["TC-001", "Создание корректной брони", "Выбрать комнату, заполнить корректные данные, нажать «Забронировать».", "Бронь появляется в таблице.", "Заполняется при тестировании.", "Passed/Failed"],
                  ["TC-002", "Проверка пересечения времени", "Создать две брони одной комнаты на пересекающееся время.", "Вторая бронь отклоняется.", "Заполняется при тестировании.", "Passed/Failed"],
                  ["TC-003", "Проверка некорректного интервала", "Ввести окончание раньше начала.", "Система показывает ошибку.", "Заполняется при тестировании.", "Passed/Failed"],
                  ["TC-004", "Проверка вместимости", "Для рабочего места указать 2 человека.", "Система показывает ошибку.", "Заполняется при тестировании.", "Passed/Failed"],
                  ["TC-005", "Сохранение цели бронирования", "Создать бронь с целью, сохранить, перезапустить приложение.", "Цель сохраняется.", "Заполняется при тестировании.", "Passed/Failed"],
              ])

    document.add_heading("7. Заключение", level=1)
    add_paragraph(document, "В результате работы создан простой программный продукт для бронирования офисных ресурсов, подготовлены техническое задание, руководство пользователя и перечень намеренно внесенных дефектов. Проект пригоден для передачи другой команде и выполнения тестирования методом «черного ящика».")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

